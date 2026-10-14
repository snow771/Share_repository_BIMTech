# -*- coding: utf-8 -*-
"""
stencilデータ変換_完成版.py
-------------------------------------------------
・フォルダ内の複数 Excel（.xlsm / .xlsx）をまとめて読込
・元データ.xlsm 系のブックから物件諸情報を抽出
・完成イメージ.xlsx と同じ列構成・近い書式で 1 つの .xlsx に出力
・マッピング YAML は自動検出（見つからない場合のみ選択）
-------------------------------------------------
"""

import re
import sys
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import yaml
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
from openpyxl.utils import get_column_letter

# ---------- GUI ----------
import tkinter as tk
from tkinter import filedialog, messagebox

root = None

# ---------- 定数 ----------
OUTPUT_COLUMNS = [
    "project_id", "物件名", "現場住所_都道府県", "現場住所_市区町村", "現場住所_町名",
    "請負金額", "契約時予定利益", "利益率", "敷地面積", "建築面積", "法廷床面積", "施工床面積",
    "施工坪単価", "工期(着工)", "工期(竣工)", "構造", "階数_地上", "階数_地下",
    "1R", "1K", "1DK", "1LDK", "2K", "2DK", "2LDK", "3K", "3DK", "3LDK",
    "4K", "4DK", "4LDK", "5K", "5DK", "5LDK", "1SK", "2SK", "１SLDK", "2SLDK", "3SLDK",
    "OR", "テナント", "店舗", "事務所", "その他", "間取", "世帯数", "形態", "ELV",
    "ｺﾝｸﾘｰﾄ", "型枠", "鉄筋（Kg/m2）", "鉄骨（Kg/m2）", "ＡＬＣ", "外壁タイル",
    "山留工法", "杭工事", "土工事", "型枠・鉄筋", "鉄骨", "ALC", "生ｺﾝ車", "ﾎﾟﾝﾌﾟ車"
]

PLAN_KEYS = [
    "1R", "1K", "1DK", "1LDK", "2K", "2DK", "2LDK", "3K", "3DK", "3LDK",
    "4K", "4DK", "4LDK", "5K", "5DK", "5LDK", "1SK", "2SK", "1SLDK", "2SLDK", "3SLDK",
]

CATEGORY_KEYS = ["OR", "テナント", "店舗", "事務所", "その他"]

REQUIRED_SHEETS = ["実行予算表書", "実行予算・引継書", "原価推移表"]

HEADER_FILL = PatternFill(fill_type="solid", start_color="ADD8E6", end_color="ADD8E6")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
HEADER_FONT = Font(name="游ゴシック", size=11, bold=True)
BODY_FONT = Font(name="游ゴシック", size=11)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="top")

COLUMN_WIDTHS = {
    "project_id": 11.73, "物件名": 62.54, "現場住所_都道府県": 21.09, "現場住所_市区町村": 21.09,
    "現場住所_町名": 21.00, "請負金額": 16.91, "契約時予定利益": 17.82, "利益率": 9.27,
    "敷地面積": 10.63, "建築面積": 13.00, "法廷床面積": 12.91, "施工床面積": 13.00,
    "施工坪単価": 12.82, "工期(着工)": 12.63, "工期(竣工)": 13.00, "構造": 25.91,
    "階数_地上": 11.45, "階数_地下": 13.00, "1R": 11.45, "1K": 13.00, "1DK": 13.00,
    "1LDK": 13.00, "2K": 13.00, "2DK": 13.00, "2LDK": 13.00, "3K": 13.00, "3DK": 13.00,
    "3LDK": 13.63, "4K": 11.45, "4DK": 13.00, "4LDK": 13.00, "5K": 13.00, "5DK": 13.00,
    "5LDK": 13.00, "1SK": 13.00, "2SK": 13.00, "１SLDK": 13.00, "2SLDK": 13.00, "3SLDK": 13.00,
    "OR": 11.45, "テナント": 13.00, "店舗": 13.00, "事務所": 13.00, "その他": 13.00,
    "間取": 21.54, "世帯数": 8.45, "形態": 7.73, "ELV": 5.73, "ｺﾝｸﾘｰﾄ": 9.00,
    "型枠": 10.09, "鉄筋（Kg/m2）": 15.63, "鉄骨（Kg/m2）": 13.00, "ＡＬＣ": 9.00,
    "外壁タイル": 11.82, "山留工法": 22.18, "杭工事": 11.18, "土工事": 8.45,
    "型枠・鉄筋": 11.82, "鉄骨": 6.27, "ALC": 7.73, "生ｺﾝ車": 8.54, "ﾎﾟﾝﾌﾟ車": 9.09,
}

# ---------- 補助関数 ----------
def to_nfk(text: Any) -> str:
    if text is None:
        return ""
    return unicodedata.normalize("NFKC", str(text)).strip()


def safe_int(value: Any, default: int = 0) -> int:
    if value in (None, ""):
        return default
    try:
        return int(float(value))
    except Exception:
        text = to_nfk(value)
        nums = re.findall(r"\d+", text)
        return int(nums[0]) if nums else default


def safe_number(value: Any) -> Optional[float]:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return value
    text = to_nfk(value).replace(",", "")
    try:
        return float(text)
    except Exception:
        return None


def normalize_plan_key(plan: str) -> str:
    plan = to_nfk(plan).upper().replace(" ", "")
    if plan == "1SLDK":
        return "１SLDK"
    return plan


def split_japanese_address(address: Any) -> Tuple[str, str, str]:
    text = to_nfk(address)
    if not text:
        return "", "", ""

    pref_match = re.match(r"^(東京都|北海道|(?:京都|大阪)府|.{2,3}県)", text)
    prefecture = pref_match.group(1) if pref_match else ""
    rest = text[len(prefecture):] if prefecture else text

    municipality = ""
    # 政令指定都市 + 区
    m = re.match(r"^(.+?市.+?区)", rest)
    if m:
        municipality = m.group(1)
    else:
        m = re.match(r"^(.+?[市区町村])", rest)
        if m:
            municipality = m.group(1)

    tail = rest[len(municipality):] if municipality else rest
    town = ""

    m = re.match(r"^(.+?\d+丁目)", tail)
    if m:
        town = m.group(1)
    else:
        m = re.match(r"^([^0-9\-－―ー]+)", tail)
        if m:
            town = m.group(1).strip()

    return prefecture, municipality, town


def parse_plan_text(plan_text: Any) -> Dict[str, int]:
    result = {k: 0 for k in PLAN_KEYS + CATEGORY_KEYS}
    text = to_nfk(plan_text)
    if not text:
        return result

    normalized = text.replace("\r", "\n").replace("\n", "、")

    # 間取り戸数
    for plan, count in re.findall(r"((?:\d+|OR)(?:R|K|DK|LDK|SK|SLDK))\s*[-－ー]?\s*(\d+)\s*戸?", normalized, flags=re.IGNORECASE):
        key = normalize_plan_key(plan)
        if key in result:
            result[key] += int(count)

    # 用途戸数（店舗・事務所・テナント等）
    for cat in CATEGORY_KEYS:
        pattern = rf"{re.escape(cat)}\s*[-－ー]?\s*(\d+)"
        for count in re.findall(pattern, normalized, flags=re.IGNORECASE):
            result[cat] += int(count)

    return result


def find_mapping_yaml(script_dir: Path, folder: Path) -> Optional[Path]:
    candidates: List[Path] = []
    for base in [script_dir, folder, Path.cwd()]:
        if not base.exists():
            continue
        candidates.extend(sorted(base.glob("*.yml")))
        candidates.extend(sorted(base.glob("*.yaml")))

    for p in candidates:
        if "mapping" in p.stem.lower() or "マッピング" in p.stem:
            return p
    return candidates[0] if candidates else None


def load_yaml_mapping(path: Path) -> Dict[str, Any]:
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def get_cell_by_mapping(wb, mapping: Dict[str, Any], section: str, field: str) -> Any:
    try:
        sec = mapping.get(section, {})
        sheet_name = sec.get("sheet")
        addr = sec.get("fields", {}).get(field)
        if not sheet_name or not addr or sheet_name not in wb.sheetnames:
            return None
        return wb[sheet_name][addr].value
    except Exception:
        return None


def extract_row_from_workbook(excel_path: Path, mapping: Dict[str, Any]) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"必要シート不足: {', '.join(missing)}")

    row = {col: None for col in OUTPUT_COLUMNS}

    # 主要項目
    row["project_id"] = get_cell_by_mapping(wb, mapping, "実行予算表書", "現場コード")
    row["物件名"] = get_cell_by_mapping(wb, mapping, "実行予算表書", "物件名")
    address = get_cell_by_mapping(wb, mapping, "実行予算表書", "現場住所")
    pref, city, town = split_japanese_address(address)
    row["現場住所_都道府県"] = pref
    row["現場住所_市区町村"] = city
    row["現場住所_町名"] = town

    for key in ["請負金額", "契約時予定利益", "利益率", "敷地面積", "建築面積", "法廷床面積", "施工床面積", "構造", "世帯数", "ｺﾝｸﾘｰﾄ", "型枠", "鉄筋（Kg/m2）", "鉄骨（Kg/m2）", "ＡＬＣ", "外壁タイル"]:
        row[key] = get_cell_by_mapping(wb, mapping, "実行予算表書", key)

    # 施工坪単価
    unit_price = get_cell_by_mapping(wb, mapping, "実行予算表書", "施工坪")
    if unit_price is None:
        contract = safe_number(row["請負金額"])
        floor_area = safe_number(row["施工床面積"])
        if contract is not None and floor_area not in (None, 0):
            tsubo = floor_area / 3.305785
            unit_price = contract / tsubo if tsubo else None
    row["施工坪単価"] = unit_price

    # 日付
    row["工期(着工)"] = get_cell_by_mapping(wb, mapping, "実行予算表書", "工期(着工､竣工)")
    row["工期(竣工)"] = get_cell_by_mapping(wb, mapping, "実行予算表書", "工期(着工､竣工)_2")

    # 階数
    row["階数_地上"] = safe_int(get_cell_by_mapping(wb, mapping, "実行予算表書", "階数_地上"), 0)
    row["階数_地下"] = safe_int(get_cell_by_mapping(wb, mapping, "実行予算表書", "階数_地下"), 0)

    # 間取関連
    plan_text = get_cell_by_mapping(wb, mapping, "実行予算表書", "間取")
    counts = parse_plan_text(plan_text)
    for key in PLAN_KEYS:
        out_key = "１SLDK" if key == "1SLDK" else key
        row[out_key] = counts.get(key, 0)
    for key in CATEGORY_KEYS:
        row[key] = counts.get(key, 0)
    row["間取"] = None

    # 原価推移表
    row["形態"] = get_cell_by_mapping(wb, mapping, "原価推移表", "形態")
    row["ELV"] = get_cell_by_mapping(wb, mapping, "原価推移表", "ELV")

    # 引継書
    for key in ["山留工法", "杭工事", "土工事", "型枠・鉄筋", "鉄骨", "ALC", "生ｺﾝ車", "ﾎﾟﾝﾌﾟ車"]:
        row[key] = get_cell_by_mapping(wb, mapping, "実行予算・引継書", key)

    return row


def apply_output_format(ws):
    # ヘッダー
    for c, header in enumerate(OUTPUT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        ws.column_dimensions[get_column_letter(c)].width = COLUMN_WIDTHS.get(header, 13)

    # データ行
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = THIN_BORDER
            cell.font = BODY_FONT

    # 表示形式
    index = {name: i + 1 for i, name in enumerate(OUTPUT_COLUMNS)}
    int_cols = ["請負金額", "契約時予定利益"]
    one_decimal_cols = ["敷地面積", "建築面積", "法廷床面積", "施工床面積", "施工坪単価"]
    date_cols = ["工期(着工)", "工期(竣工)"]

    for row_no in range(2, ws.max_row + 1):
        for name in int_cols:
            ws.cell(row_no, index[name]).number_format = "#,##0"
        for name in one_decimal_cols:
            ws.cell(row_no, index[name]).number_format = "#,##0.0"
        ws.cell(row_no, index["利益率"]).number_format = "0.00000_);[Red]\\(0.00000\\)"
        for name in date_cols:
            ws.cell(row_no, index[name]).number_format = "mm-dd-yy"

    ws.freeze_panes = "AT2"


def collect_excel_files(folder: Path) -> List[Path]:
    files: List[Path] = []
    for ext in ("*.xlsm", "*.xlsx"):
        files.extend(folder.rglob(ext))
    result = []
    for p in sorted(files):
        name = p.name
        if name.startswith("~$"):
            continue
        if p.suffix.lower() not in (".xlsm", ".xlsx"):
            continue
        result.append(p)
    return result


def choose_inputs() -> Tuple[Path, Path, Path]:
    global root
    if root is None:
        root = tk.Tk()
        root.withdraw()
    folder = filedialog.askdirectory(title="元データ Excel が入っているフォルダを選択してください")
    if not folder:
        messagebox.showinfo("中止", "フォルダが選択されませんでした。")
        sys.exit()
    folder_path = Path(folder)

    script_dir = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
    yaml_path = find_mapping_yaml(script_dir, folder_path)
    if yaml_path is None:
        selected = filedialog.askopenfilename(
            title="マッピング YAML を選択してください",
            filetypes=[("YAML", "*.yaml *.yml"), ("All files", "*.*")],
        )
        if not selected:
            messagebox.showinfo("中止", "YAML が選択されませんでした。")
            sys.exit()
        yaml_path = Path(selected)

    output_path = filedialog.asksaveasfilename(
        title="出力先ファイル名を指定してください",
        defaultextension=".xlsx",
        filetypes=[("Excel", "*.xlsx")],
        initialfile="物件諸情報一覧.xlsx",
    )
    if not output_path:
        messagebox.showinfo("中止", "保存先が選択されませんでした。")
        sys.exit()

    return folder_path, yaml_path, Path(output_path)


def main():
    folder_path, yaml_path, output_path = choose_inputs()
    mapping = load_yaml_mapping(yaml_path)
    excel_files = collect_excel_files(folder_path)

    if not excel_files:
        messagebox.showerror("エラー", "対象となる Excel ファイルが見つかりませんでした。")
        return

    rows: List[Dict[str, Any]] = []
    skipped: List[str] = []

    for excel_path in excel_files:
        try:
            row = extract_row_from_workbook(excel_path, mapping)
            if not row.get("project_id") and not row.get("物件名"):
                skipped.append(f"{excel_path.name} : 抽出対象データなし")
                continue
            rows.append(row)
        except Exception as e:
            skipped.append(f"{excel_path.name} : {e}")

    if not rows:
        messagebox.showerror("エラー", "処理できる Excel ファイルがありませんでした。")
        return

    # 並び順を project_id → 物件名
    rows.sort(key=lambda x: (str(x.get("project_id") or ""), str(x.get("物件名") or "")))

    wb = Workbook()
    ws = wb.active
    ws.title = output_path.stem[:31]

    apply_output_format(ws)

    for r_idx, row_data in enumerate(rows, start=2):
        for c_idx, header in enumerate(OUTPUT_COLUMNS, start=1):
            ws.cell(row=r_idx, column=c_idx, value=row_data.get(header))

    # 罫線・フォント・数値書式をデータ反映後に再適用
    apply_output_format(ws)
    wb.save(output_path)

    if skipped:
        preview = "\n".join(skipped[:15])
        suffix = "\n..." if len(skipped) > 15 else ""
        messagebox.showwarning(
            "完了（スキップあり）",
            f"出力完了:\n{output_path}\n\nスキップ件数: {len(skipped)}\n{preview}{suffix}"
        )
    else:
        messagebox.showinfo("完了", f"出力完了:\n{output_path}")


if __name__ == "__main__":
    main()
