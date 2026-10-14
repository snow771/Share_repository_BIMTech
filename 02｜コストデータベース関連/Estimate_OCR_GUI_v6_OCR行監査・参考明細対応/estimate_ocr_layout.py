"""
協力業者見積PDFを Azure AI Document Intelligence の prebuilt-layout で解析し、
生JSONと確認用Excelを出力する検証スクリプト。

実行例:
    python estimate_ocr_layout.py "C:\OCR_Test\見積書サンプル.pdf"

必要な環境変数:
    DI_ENDPOINT
    DI_KEY
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from azure.ai.documentintelligence import DocumentIntelligenceClient
from azure.core.credentials import AzureKeyCredential
from azure.core.exceptions import AzureError
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="見積PDFをAzure Document Intelligence Layoutモデルで解析します。"
    )
    parser.add_argument(
        "pdf_path",
        help="解析するPDFファイルのパス。例: C:\\OCR_Test\\見積書.pdf",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="出力先フォルダ。未指定時はPDFと同じ場所に OCR_出力 フォルダを作成します。",
    )
    return parser.parse_args()


def get_client() -> DocumentIntelligenceClient:
    endpoint = os.getenv("DI_ENDPOINT")
    key = os.getenv("DI_KEY")

    missing = []
    if not endpoint:
        missing.append("DI_ENDPOINT")
    if not key:
        missing.append("DI_KEY")

    if missing:
        names = ", ".join(missing)
        raise RuntimeError(
            f"環境変数 {names} が見つかりません。\n"
            "PowerShellで setx を実行した後、新しいPowerShellを開いてから再実行してください。"
        )

    return DocumentIntelligenceClient(
        endpoint=endpoint,
        credential=AzureKeyCredential(key),
    )


def analyze_pdf(client: DocumentIntelligenceClient, pdf_path: Path) -> Any:
    with pdf_path.open("rb") as pdf_file:
        poller = client.begin_analyze_document(
            "prebuilt-layout",
            body=pdf_file,
        )
    return poller.result()


def safe_value(value: Any) -> Any:
    """Azure SDKのオブジェクトをJSONに書ける値へ寄せる。"""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, list):
        return [safe_value(v) for v in value]
    if isinstance(value, dict):
        return {str(k): safe_value(v) for k, v in value.items()}
    return str(value)


def write_raw_json(result: Any, json_path: Path) -> None:
    if hasattr(result, "as_dict"):
        data = result.as_dict()
    else:
        data = safe_value(result)

    with json_path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2, default=safe_value)


def get_page_numbers(bounding_regions: Any) -> str:
    if not bounding_regions:
        return ""
    numbers = []
    for region in bounding_regions:
        page_number = getattr(region, "page_number", None)
        if page_number is not None:
            numbers.append(str(page_number))
    return ", ".join(sorted(set(numbers), key=lambda x: int(x)))


def set_header_style(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def fit_columns(ws, max_width: int = 50) -> None:
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        longest = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            longest = max(longest, len(value))
        ws.column_dimensions[column_letter].width = min(max(longest + 2, 10), max_width)


def write_excel(result: Any, excel_path: Path, source_pdf: Path) -> None:
    wb = Workbook()

    # 概要シート
    ws = wb.active
    ws.title = "概要"
    ws.append(["項目", "値"])
    ws.append(["元PDF", source_pdf.name])
    ws.append(["作成日時", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append(["モデル", getattr(result, "model_id", "")])
    ws.append(["APIバージョン", getattr(result, "api_version", "")])
    ws.append(["ページ数", len(result.pages or [])])
    ws.append(["検出表数", len(result.tables or [])])
    ws.append(["段落数", len(result.paragraphs or [])])
    set_header_style(ws)
    fit_columns(ws)

    # OCR行の一覧
    ws_lines = wb.create_sheet("OCR行一覧")
    ws_lines.append(["ページ", "行番号", "文字列"])
    for page in result.pages or []:
        for line_index, line in enumerate(page.lines or [], start=1):
            ws_lines.append([
                getattr(page, "page_number", ""),
                line_index,
                getattr(line, "content", ""),
            ])
    set_header_style(ws_lines)
    ws_lines.column_dimensions["A"].width = 10
    ws_lines.column_dimensions["B"].width = 10
    ws_lines.column_dimensions["C"].width = 80
    for row in ws_lines.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    # 各表をシートへ出力
    tables = result.tables or []
    for table_index, table in enumerate(tables, start=1):
        sheet_name = f"表_{table_index:02d}"
        ws_table = wb.create_sheet(sheet_name)

        ws_table.append([
            f"表番号: {table_index}",
            f"ページ: {get_page_numbers(getattr(table, 'bounding_regions', None))}",
            f"行数: {table.row_count}",
            f"列数: {table.column_count}",
        ])
        ws_table.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(table.column_count, 4))
        ws_table["A1"].font = Font(bold=True)
        ws_table["A1"].fill = PatternFill("solid", fgColor="D9EAF7")

        # 表のセルを行・列位置に配置
        for _ in range(table.row_count):
            ws_table.append([""] * table.column_count)

        for cell in table.cells:
            target_row = cell.row_index + 2
            target_col = cell.column_index + 1
            ws_table.cell(row=target_row, column=target_col, value=cell.content)

            kind = str(getattr(cell, "kind", "") or "")
            if "columnHeader" in kind or "rowHeader" in kind:
                ws_table.cell(row=target_row, column=target_col).font = Font(bold=True)
                ws_table.cell(row=target_row, column=target_col).fill = PatternFill(
                    "solid", fgColor="E2F0D9"
                )

            ws_table.cell(row=target_row, column=target_col).alignment = Alignment(
                wrap_text=True, vertical="top"
            )

        ws_table.freeze_panes = "A2"
        fit_columns(ws_table, max_width=35)

    wb.save(excel_path)


def main() -> int:
    args = parse_args()
    pdf_path = Path(args.pdf_path).expanduser().resolve()

    if not pdf_path.exists():
        print(f"エラー: PDFが見つかりません: {pdf_path}", file=sys.stderr)
        return 1

    if pdf_path.suffix.lower() != ".pdf":
        print("エラー: 現在はPDFファイルのみ対応しています。", file=sys.stderr)
        return 1

    output_dir = (
        Path(args.output_dir).expanduser().resolve()
        if args.output_dir
        else pdf_path.parent / "OCR_出力"
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = pdf_path.stem
    json_path = output_dir / f"{stem}_{timestamp}_layout_raw.json"
    excel_path = output_dir / f"{stem}_{timestamp}_layout_check.xlsx"

    try:
        client = get_client()
        print("Azure Document IntelligenceへPDFを送信しています...")
        result = analyze_pdf(client, pdf_path)

        print("JSONを書き出しています...")
        write_raw_json(result, json_path)

        print("確認用Excelを書き出しています...")
        write_excel(result, excel_path, pdf_path)

    except RuntimeError as exc:
        print(f"設定エラー: {exc}", file=sys.stderr)
        return 1
    except AzureError as exc:
        print("Azure Document Intelligenceの呼び出しでエラーが発生しました。", file=sys.stderr)
        print(f"詳細: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print("予期しないエラーが発生しました。", file=sys.stderr)
        print(f"詳細: {exc}", file=sys.stderr)
        return 1

    print("")
    print("解析が完了しました。")
    print(f"JSON : {json_path}")
    print(f"Excel: {excel_path}")
    print(f"ページ数: {len(result.pages or [])}")
    print(f"検出表数: {len(result.tables or [])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
