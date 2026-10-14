"""
Azure Document Intelligence F0向け：PDFを指定ページ数ごとに分割してLayout OCRを実行し、
PDF全体を1つのLayout JSONと確認用Excelへ統合する。

例:
python .\estimate_ocr_batch_f0.py ".\見積.pdf" --output-dir ".\OCR結果" --chunk-pages 2

注意:
- F0の1回あたりの取得ページ上限を回避するための分割実行。
- 元PDFは変更しない。一時分割PDFは処理後に削除する。
- Azureへは分割数分のリクエストを送る。
"""

from __future__ import annotations

import argparse
import copy
import json
import os
import re
import sys
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from azure.ai.documentintelligence import DocumentIntelligenceClient
from azure.core.credentials import AzureKeyCredential
from azure.core.exceptions import AzureError
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader, PdfWriter


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="PDFを2ページ単位などに分割し、Azure Layout OCR結果を1つへ統合します。"
    )
    parser.add_argument("pdf_path", help="入力PDF")
    parser.add_argument("--output-dir", required=True, help="出力先フォルダ")
    parser.add_argument(
        "--chunk-pages",
        type=int,
        default=2,
        help="1回のOCRへ送るページ数。F0運用では通常2（既定値）。",
    )
    parser.add_argument(
        "--keep-chunks",
        action="store_true",
        help="一時分割PDFを出力先の _chunks フォルダへ残す。",
    )
    return parser.parse_args()


def get_client() -> DocumentIntelligenceClient:
    endpoint = os.getenv("DI_ENDPOINT", "").strip()
    key = os.getenv("DI_KEY", "").strip()
    if not endpoint or not key:
        raise RuntimeError(
            "DI_ENDPOINT または DI_KEY が設定されていません。"
            "PowerShellで環境変数を設定し、新しいPowerShellまたはGUIを起動してください。"
        )
    return DocumentIntelligenceClient(
        endpoint=endpoint,
        credential=AzureKeyCredential(key),
    )


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def analyze_bytes(client: DocumentIntelligenceClient, pdf_bytes: bytes) -> dict[str, Any]:
    poller = client.begin_analyze_document(
        "prebuilt-layout",
        body=BytesIO(pdf_bytes),
    )
    return poller.result().as_dict()


def recursively_adjust_page_and_offsets(
    value: Any,
    page_offset: int,
    content_offset: int,
) -> Any:
    """
    Layout結果内の pageNumber と span.offset を統合後の値へ補正する。
    構造化スクリプトは主に tables / cells / boundingRegions を使うが、
    生JSONとしても追跡しやすいようにオフセットも補正する。
    """
    if isinstance(value, list):
        return [
            recursively_adjust_page_and_offsets(v, page_offset, content_offset)
            for v in value
        ]
    if isinstance(value, dict):
        result = {}
        for key, item in value.items():
            if key == "pageNumber" and isinstance(item, int):
                result[key] = item + page_offset
            elif key == "offset" and isinstance(item, int):
                result[key] = item + content_offset
            else:
                result[key] = recursively_adjust_page_and_offsets(
                    item, page_offset, content_offset
                )
        return result
    return copy.deepcopy(value)


def merge_layout_results(
    chunk_results: list[dict[str, Any]],
    source_pdf_name: str,
    total_pages: int,
    chunk_pages: int,
) -> dict[str, Any]:
    if not chunk_results:
        raise ValueError("統合対象のOCR結果がありません。")

    first = copy.deepcopy(chunk_results[0]["result"])
    merged: dict[str, Any] = {
        key: copy.deepcopy(value)
        for key, value in first.items()
        if key not in {
            "content",
            "pages",
            "tables",
            "paragraphs",
            "sections",
            "styles",
            "keyValuePairs",
            "languages",
            "figures",
            "documents",
        }
    }

    list_keys = [
        "pages",
        "tables",
        "paragraphs",
        "sections",
        "styles",
        "keyValuePairs",
        "languages",
        "figures",
        "documents",
    ]
    for key in list_keys:
        merged[key] = []

    content_parts: list[str] = []
    content_offset = 0

    for chunk in chunk_results:
        result = chunk["result"]
        page_offset = chunk["start_page"] - 1
        chunk_content = result.get("content", "")

        for key in list_keys:
            if result.get(key):
                adjusted = recursively_adjust_page_and_offsets(
                    result[key],
                    page_offset=page_offset,
                    content_offset=content_offset,
                )
                merged[key].extend(adjusted)

        content_parts.append(chunk_content)
        content_offset += len(chunk_content) + 1

    merged["content"] = "\n".join(content_parts)
    merged["batch_info"] = {
        "source_pdf": source_pdf_name,
        "source_pdf_page_count": total_pages,
        "ocr_page_count": len(merged.get("pages", [])),
        "chunk_pages": chunk_pages,
        "chunks": [
            {
                "chunk_index": c["chunk_index"],
                "start_page": c["start_page"],
                "end_page": c["end_page"],
                "ocr_page_count": len(c["result"].get("pages", [])),
            }
            for c in chunk_results
        ],
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    return merged


def table_matrix(table: dict[str, Any]) -> list[list[str]]:
    row_count = int(table.get("rowCount", 0))
    col_count = int(table.get("columnCount", 0))
    rows = [["" for _ in range(col_count)] for _ in range(row_count)]
    for cell in table.get("cells", []):
        r = int(cell.get("rowIndex", 0))
        c = int(cell.get("columnIndex", 0))
        if 0 <= r < row_count and 0 <= c < col_count:
            rows[r][c] = clean_text(cell.get("content", ""))
    return rows


def table_page(table: dict[str, Any]) -> int | None:
    regions = table.get("boundingRegions") or []
    return regions[0].get("pageNumber") if regions else None


def set_header_style(ws, row_number: int) -> None:
    fill = PatternFill("solid", fgColor="2F7D62")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row_number]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def fit_columns(ws, max_width: int = 40) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = 8
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            longest = max((len(line) for line in value.splitlines()), default=0)
            width = max(width, min(longest + 2, max_width))
        ws.column_dimensions[letter].width = width


def write_check_excel(merged: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Overview"

    info = merged.get("batch_info", {})
    summary.append(["項目", "値"])
    summary.append(["元PDF", info.get("source_pdf", "")])
    summary.append(["元PDFページ数", info.get("source_pdf_page_count", "")])
    summary.append(["Azure取得ページ数", info.get("ocr_page_count", "")])
    summary.append(["分割単位", f"{info.get('chunk_pages', '')}ページ"])
    summary.append(["検出表数", len(merged.get("tables", []))])
    set_header_style(summary, 1)
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    fit_columns(summary)

    lines = wb.create_sheet("OCR line list")
    lines.append(["PDFページ", "OCR行"])
    for page in merged.get("pages", []):
        page_no = page.get("pageNumber")
        for line in page.get("lines", []):
            lines.append([page_no, line.get("content", "")])
    set_header_style(lines, 1)
    lines.freeze_panes = "A2"
    fit_columns(lines, max_width=80)

    for table_index, table in enumerate(merged.get("tables", []), start=1):
        ws = wb.create_sheet(f"Table_{table_index}")
        ws.append([f"PDFページ: {table_page(table) or ''}"])
        rows = table_matrix(table)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F5E4B")
            cell.font = Font(color="FFFFFF", bold=True)
        if rows:
            set_header_style(ws, 2)
        ws.freeze_panes = "A3"
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        fit_columns(ws)

    wb.save(path)


def main() -> int:
    args = parse_args()
    pdf_path = Path(args.pdf_path).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()

    if not pdf_path.exists():
        print(f"エラー: PDFが見つかりません: {pdf_path}", file=sys.stderr)
        return 1
    if pdf_path.suffix.lower() != ".pdf":
        print("エラー: PDFファイルを指定してください。", file=sys.stderr)
        return 1
    if args.chunk_pages < 1:
        print("エラー: --chunk-pages は1以上にしてください。", file=sys.stderr)
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(str(pdf_path))
    total_pages = len(reader.pages)
    if total_pages == 0:
        print("エラー: PDFページがありません。", file=sys.stderr)
        return 1

    client = get_client()
    chunk_results: list[dict[str, Any]] = []

    keep_dir = output_dir / "_chunks"
    if args.keep_chunks:
        keep_dir.mkdir(exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="estimate_ocr_chunks_") as temp_dir_name:
        temp_dir = Path(temp_dir_name)

        chunk_index = 0
        for start_zero in range(0, total_pages, args.chunk_pages):
            chunk_index += 1
            end_zero = min(start_zero + args.chunk_pages, total_pages)
            start_page = start_zero + 1
            end_page = end_zero

            writer = PdfWriter()
            for page_index in range(start_zero, end_zero):
                writer.add_page(reader.pages[page_index])

            chunk_path = temp_dir / f"{pdf_path.stem}_p{start_page:03d}-{end_page:03d}.pdf"
            with chunk_path.open("wb") as f:
                writer.write(f)

            if args.keep_chunks:
                target = keep_dir / chunk_path.name
                target.write_bytes(chunk_path.read_bytes())

            print(
                f"OCR処理中: {chunk_index}ブロック目 "
                f"（元PDF {start_page}-{end_page}ページ / 全{total_pages}ページ）"
            )
            try:
                result = analyze_bytes(client, chunk_path.read_bytes())
            except AzureError as exc:
                print(
                    f"エラー: {start_page}-{end_page}ページのOCRに失敗しました: {exc}",
                    file=sys.stderr,
                )
                return 1

            actual_pages = len(result.get("pages", []))
            if actual_pages != (end_page - start_page + 1):
                print(
                    f"警告: {start_page}-{end_page}ページを送信しましたが、"
                    f"Azureは{actual_pages}ページを返しました。",
                    file=sys.stderr,
                )

            chunk_results.append(
                {
                    "chunk_index": chunk_index,
                    "start_page": start_page,
                    "end_page": end_page,
                    "result": result,
                }
            )

    merged = merge_layout_results(
        chunk_results,
        source_pdf_name=pdf_path.name,
        total_pages=total_pages,
        chunk_pages=args.chunk_pages,
    )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    raw_json_path = output_dir / f"{pdf_path.stem}_{timestamp}_layout_raw_full.json"
    check_xlsx_path = output_dir / f"{pdf_path.stem}_{timestamp}_layout_check_full.xlsx"
    raw_json_path.write_text(
        json.dumps(merged, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    write_check_excel(merged, check_xlsx_path)

    print("全ページOCR・統合が完了しました。")
    print(f"元PDFページ数     : {total_pages}")
    print(f"Azure取得ページ数 : {len(merged.get('pages', []))}")
    print(f"統合JSON          : {raw_json_path}")
    print(f"確認用Excel       : {check_xlsx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
